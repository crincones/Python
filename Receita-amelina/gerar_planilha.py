# -*- coding: utf-8 -*-
"""Gera a planilha de custos das receitas (pronta para importar no Google Planilhas)."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- estilos ----
AZUL      = "1F3864"
AZUL_CLR  = "D9E2F3"
AMARELO   = "FFF2CC"   # campo editável
CINZA     = "F2F2F2"
VERDE     = "E2EFDA"
LARANJA   = "FCE4D6"

f_titulo  = Font(name="Calibri", size=16, bold=True, color=AZUL)
f_sub     = Font(name="Calibri", size=10, italic=True, color="595959")
f_head    = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
f_secao   = Font(name="Calibri", size=11, bold=True, color=AZUL)
f_bold    = Font(name="Calibri", size=11, bold=True)
f_total   = Font(name="Calibri", size=12, bold=True, color=AZUL)

p_head    = PatternFill("solid", fgColor=AZUL)
p_edit    = PatternFill("solid", fgColor=AMARELO)
p_secao   = PatternFill("solid", fgColor=AZUL_CLR)
p_calc    = PatternFill("solid", fgColor=CINZA)
p_result  = PatternFill("solid", fgColor=VERDE)
p_destaq  = PatternFill("solid", fgColor=LARANJA)

_fino = Side(style="thin", color="BFBFBF")
b_all = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)

MOEDA   = 'R$ #,##0.00'
MOEDA4  = 'R$ #,##0.0000'
PCT     = '0.0%'
NUM     = '#,##0.###'


def cel(ws, ref, valor=None, fill=None, font=None, fmt=None, align=None, borda=True, wrap=False):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    if fill:
        c.fill = fill
    if font:
        c.font = font
    if fmt:
        c.number_format = fmt
    if align or wrap:
        c.alignment = Alignment(horizontal=align or "general", vertical="center", wrap_text=wrap)
    if borda:
        c.border = b_all
    return c


wb = Workbook()

# ============================================================== INSUMOS ======
ws = wb.active
ws.title = "1. Insumos"
ws.sheet_view.showGridLines = False

ws["A1"] = "TABELA DE INSUMOS - preços de compra"
ws["A1"].font = f_titulo
ws["A2"] = ("Altere aqui o preço e o tamanho da embalagem que você compra. "
            "Todas as receitas recalculam sozinhas.")
ws["A2"].font = f_sub
ws["A3"] = "AMARELO = você edita     |     CINZA = calculado automaticamente (não digite)"
ws["A3"].font = Font(size=10, bold=True, color="7F6000")

cabs = ["Insumo", "Preço da embalagem (R$)", "Qtd. na embalagem", "Unidade", "Custo por unidade (R$)"]
for i, h in enumerate(cabs, start=1):
    cel(ws, get_column_letter(i) + "5", h, p_head, f_head, align="center", wrap=True)

insumos = [
    ("Café solúvel",   11.24,   50, "g"),
    ("Açúcar",          7.99, 2000, "g"),
    ("Cacau 50%",      35.00,  200, "g"),
    ("Canela em pó",    6.99,   20, "g"),
    ("Água",            0.00, 1000, "ml"),
    ("Pote de vidro",   3.45,    1, "un"),
    ("Adesivo frente",  0.44,    1, "un"),
    ("Adesivo verso",   0.31,    1, "un"),
]
LIN0 = 6
for k, (nome, preco, qtd, un) in enumerate(insumos):
    r = LIN0 + k
    cel(ws, "A%d" % r, nome, font=f_bold)
    cel(ws, "B%d" % r, preco, p_edit, fmt=MOEDA)
    cel(ws, "C%d" % r, qtd, p_edit, fmt=NUM, align="center")
    cel(ws, "D%d" % r, un, p_edit, align="center")
    cel(ws, "E%d" % r, "=IF(C{0}=0,0,B{0}/C{0})".format(r), p_calc, fmt=MOEDA4)
LIN1 = LIN0 + len(insumos) - 1
FAIXA = "'1. Insumos'!$A${0}:$E${1}".format(LIN0, LIN1)

ws["A%d" % (LIN1 + 2)] = ("A água não entra no custo (preço R$ 0,00). "
                          "Se quiser considerá-la, basta preencher o preço acima.")
ws["A%d" % (LIN1 + 2)].font = f_sub

for col, w in zip("ABCDE", (24, 24, 20, 12, 22)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A6"


# ============================================================== RECEITAS =====
def aba_receita(titulo, nome_aba, ingredientes, cor_aba):
    ws = wb.create_sheet(nome_aba)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = cor_aba

    ws["A1"] = titulo
    ws["A1"].font = f_titulo
    ws["A2"] = "AMARELO = você edita     |     CINZA / VERDE / LARANJA = calculado automaticamente"
    ws["A2"].font = Font(size=10, bold=True, color="7F6000")

    cel(ws, "A4", "Rendimento do lote (potes)", p_secao, f_secao)
    cel(ws, "B4", 4, p_edit, f_bold, fmt=NUM, align="center")
    cel(ws, "C4", "<- quantos potes saem de uma receita", font=f_sub, borda=False)
    REND = "$B$4"

    # ---- ingredientes
    r = 6
    cel(ws, "A%d" % r, "INGREDIENTES (quantidade para o lote inteiro)", p_secao, f_secao)
    for c in "BCDE":
        cel(ws, "%s%d" % (c, r), None, p_secao)
    r += 1
    for i, h in enumerate(["Ingrediente", "Qtd. no lote", "Unidade",
                           "Custo unitário (R$)", "Custo no lote (R$)"], 1):
        cel(ws, "%s%d" % (get_column_letter(i), r), h, p_head, f_head, align="center", wrap=True)
    ini_ing = r + 1
    for nome, qtd, un in ingredientes:
        r += 1
        cel(ws, "A%d" % r, nome)
        cel(ws, "B%d" % r, qtd, p_edit, fmt=NUM, align="center")
        cel(ws, "C%d" % r, un, p_calc, align="center")
        cel(ws, "D%d" % r, '=IFERROR(VLOOKUP($A{0},{1},5,FALSE),0)'.format(r, FAIXA), p_calc, fmt=MOEDA4)
        cel(ws, "E%d" % r, "=B{0}*D{0}".format(r), p_calc, fmt=MOEDA)
    fim_ing = r
    r += 1
    cel(ws, "A%d" % r, "Subtotal ingredientes (lote)", p_calc, f_bold)
    for c in "BCD":
        cel(ws, "%s%d" % (c, r), None, p_calc)
    SUB_ING = "$E$%d" % r
    cel(ws, "E%d" % r, "=SUM(E{0}:E{1})".format(ini_ing, fim_ing), p_calc, f_bold, fmt=MOEDA)

    # ---- embalagem
    r += 2
    cel(ws, "A%d" % r, "EMBALAGEM E RÓTULOS (por pote)", p_secao, f_secao)
    for c in "BCDE":
        cel(ws, "%s%d" % (c, r), None, p_secao)
    r += 1
    for i, h in enumerate(["Item", "Qtd. por pote", "Unidade",
                           "Custo unitário (R$)", "Custo no lote (R$)"], 1):
        cel(ws, "%s%d" % (get_column_letter(i), r), h, p_head, f_head, align="center", wrap=True)
    ini_emb = r + 1
    for nome in ("Pote de vidro", "Adesivo frente", "Adesivo verso"):
        r += 1
        cel(ws, "A%d" % r, nome)
        cel(ws, "B%d" % r, 1, p_edit, fmt=NUM, align="center")
        cel(ws, "C%d" % r, "un", p_calc, align="center")
        cel(ws, "D%d" % r, '=IFERROR(VLOOKUP($A{0},{1},5,FALSE),0)'.format(r, FAIXA), p_calc, fmt=MOEDA4)
        cel(ws, "E%d" % r, "=B{0}*D{0}*{1}".format(r, REND), p_calc, fmt=MOEDA)
    fim_emb = r
    r += 1
    cel(ws, "A%d" % r, "Subtotal embalagem (lote)", p_calc, f_bold)
    for c in "BCD":
        cel(ws, "%s%d" % (c, r), None, p_calc)
    SUB_EMB = "$E$%d" % r
    cel(ws, "E%d" % r, "=SUM(E{0}:E{1})".format(ini_emb, fim_emb), p_calc, f_bold, fmt=MOEDA)

    # ---- outros custos
    r += 2
    cel(ws, "A%d" % r, "OUTROS CUSTOS DO LOTE (opcional)", p_secao, f_secao)
    for c in "BCDE":
        cel(ws, "%s%d" % (c, r), None, p_secao)
    r += 1
    for i, h in enumerate(["Descrição", "", "", "", "Valor no lote (R$)"], 1):
        cel(ws, "%s%d" % (get_column_letter(i), r), h, p_head, f_head, align="center", wrap=True)
    ini_out = r + 1
    for nome in ("Gás / energia", "Mão de obra", "Outros"):
        r += 1
        cel(ws, "A%d" % r, nome, p_edit)
        for c in "BCD":
            cel(ws, "%s%d" % (c, r), None, p_calc)
        cel(ws, "E%d" % r, 0, p_edit, fmt=MOEDA)
    fim_out = r
    r += 1
    cel(ws, "A%d" % r, "Subtotal outros custos (lote)", p_calc, f_bold)
    for c in "BCD":
        cel(ws, "%s%d" % (c, r), None, p_calc)
    SUB_OUT = "$E$%d" % r
    cel(ws, "E%d" % r, "=SUM(E{0}:E{1})".format(ini_out, fim_out), p_calc, f_bold, fmt=MOEDA)

    # ---- resultado
    r += 2
    cel(ws, "A%d" % r, "RESULTADO", p_secao, f_secao)
    for c in "BCDE":
        cel(ws, "%s%d" % (c, r), None, p_secao)

    estado = {"r": r}

    def linha(rot, formula, fill, fmt, font=f_bold, nota=""):
        estado["r"] += 1
        rr = estado["r"]
        cel(ws, "A%d" % rr, rot, fill, font)
        for c in "BCD":
            cel(ws, "%s%d" % (c, rr), None, fill)
        cel(ws, "E%d" % rr, formula, fill, font, fmt=fmt)
        if nota:
            cel(ws, "F%d" % rr, nota, font=f_sub, borda=False)
        return "$E$%d" % rr

    CUSTO_LOTE = linha("CUSTO TOTAL DO LOTE",
                       "={0}+{1}+{2}".format(SUB_ING, SUB_EMB, SUB_OUT),
                       p_result, MOEDA, f_total)
    CUSTO_POTE = linha("CUSTO POR POTE",
                       "=IF({0}=0,0,{1}/{0})".format(REND, CUSTO_LOTE),
                       p_result, MOEDA, f_total)

    estado["r"] += 1
    rr = estado["r"]
    cel(ws, "A%d" % rr, "PREÇO DE VENDA POR POTE", p_edit, f_total)
    for c in "BCD":
        cel(ws, "%s%d" % (c, rr), None, p_edit)
    cel(ws, "E%d" % rr, 25.00, p_edit, f_total, fmt=MOEDA)
    cel(ws, "F%d" % rr, "<- digite aqui o preço que você cobra", font=f_sub, borda=False)
    PREÇO = "$E$%d" % rr

    LUCRO_POTE = linha("Lucro por pote", "={0}-{1}".format(PREÇO, CUSTO_POTE),
                       p_destaq, MOEDA, f_total)
    linha("Faturamento do lote", "={0}*{1}".format(PREÇO, REND), p_calc, MOEDA)
    linha("Lucro do lote", "={0}*{1}".format(LUCRO_POTE, REND), p_destaq, MOEDA, f_total)
    linha("Margem de lucro (%)", "=IF({0}=0,0,{1}/{0})".format(PREÇO, LUCRO_POTE),
          p_destaq, PCT, f_total, "lucro sobre o preço de venda")
    linha("Markup (multiplicador)", "=IF({0}=0,0,{1}/{0})".format(CUSTO_POTE, PREÇO),
          p_calc, '0.00"x"', f_bold, "quantas vezes o custo")

    for col, w in zip("ABCDEF", (34, 16, 12, 20, 22, 38)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    return (ws.title, REND, CUSTO_LOTE, CUSTO_POTE, PREÇO, LUCRO_POTE)


cafe = aba_receita(
    "RECEITA 1 - CAFÉ (tradicional)", "2. Café",
    [("Café solúvel", 50, "g"), ("Açúcar", 250, "g"), ("Água", 150, "ml")],
    "C00000")

capp = aba_receita(
    "RECEITA 2 - CAPPUCCINO", "3. Cappuccino",
    [("Café solúvel", 50, "g"), ("Açúcar", 250, "g"), ("Água", 150, "ml"),
     ("Cacau 50%", 36, "g"), ("Canela em pó", 2, "g")],
    "833C00")

# ============================================================== RESUMO =======
ws = wb.create_sheet("4. Resumo", 0)
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "1F3864"

ws["A1"] = "RESUMO - custo e lucro das receitas"
ws["A1"].font = f_titulo
ws["A2"] = ("Esta aba só mostra resultados. Para mudar preços vá em '1. Insumos'; "
            "para mudar quantidades e preço de venda vá nas abas das receitas.")
ws["A2"].font = f_sub

for i, h in enumerate(["Indicador", "Receita 1 - Café", "Receita 2 - Cappuccino"], 1):
    cel(ws, "%s4" % get_column_letter(i), h, p_head, f_head, align="center", wrap=True)

linhas = [
    ("Rendimento (potes por lote)",  1, NUM,   p_calc,   f_bold),
    ("Custo total do lote",          2, MOEDA, p_calc,   f_bold),
    ("CUSTO POR POTE",               3, MOEDA, p_result, f_total),
    ("Preço de venda por pote",      4, MOEDA, p_calc,   f_bold),
    ("LUCRO POR POTE",               5, MOEDA, p_destaq, f_total),
]

r = 4
for rot, idx, fmt, fill, font in linhas:
    r += 1
    cel(ws, "A%d" % r, rot, p_secao, f_secao)
    for col, rec in (("B", cafe), ("C", capp)):
        cel(ws, "%s%d" % (col, r), "='{0}'!{1}".format(rec[0], rec[idx]),
            fill, font, fmt=fmt, align="center")

r += 1
cel(ws, "A%d" % r, "Lucro do lote", p_destaq, f_secao)
for col, rec in (("B", cafe), ("C", capp)):
    cel(ws, "%s%d" % (col, r),
        "='{0}'!{1}*'{0}'!{2}".format(rec[0], rec[5], rec[1]),
        p_destaq, f_total, fmt=MOEDA, align="center")

r += 1
cel(ws, "A%d" % r, "Margem de lucro", p_destaq, f_secao)
for col, rec in (("B", cafe), ("C", capp)):
    cel(ws, "%s%d" % (col, r),
        "=IF('{0}'!{1}=0,0,'{0}'!{2}/'{0}'!{1})".format(rec[0], rec[4], rec[5]),
        p_destaq, f_total, fmt=PCT, align="center")

# ---- simulador
r += 2
cel(ws, "A%d" % r, "SIMULADOR DE PRODUÇÃO", p_secao, f_secao)
for c in "BC":
    cel(ws, "%s%d" % (c, r), None, p_secao)
r += 1
cel(ws, "A%d" % r, "Quantos lotes você vai produzir?", p_head, f_head)
cel(ws, "B%d" % r, 1, p_edit, f_bold, fmt=NUM, align="center")
cel(ws, "C%d" % r, 1, p_edit, f_bold, fmt=NUM, align="center")
LOTES = r
cel(ws, "D%d" % r, "<- edite estes dois números (amarelo)", font=f_sub, borda=False)

sim = [
    ("Potes produzidos",  "='{0}'!{1}*{2}$%d" % LOTES, (1,),       NUM,   p_calc),
    ("Custo total",       "='{0}'!{1}*{2}$%d" % LOTES, (2,),       MOEDA, p_calc),
    ("Faturamento total", "='{0}'!{1}*'{0}'!{3}*{2}$%d" % LOTES, (4, 1), MOEDA, p_calc),
    ("Lucro total",       "='{0}'!{1}*'{0}'!{3}*{2}$%d" % LOTES, (5, 1), MOEDA, p_destaq),
]
for rot, tpl, idxs, fmt, fill in sim:
    r += 1
    cel(ws, "A%d" % r, rot, p_secao, f_secao)
    for col, rec in (("B", cafe), ("C", capp)):
        if len(idxs) == 1:
            f = tpl.format(rec[0], rec[idxs[0]], col)
        else:
            f = tpl.format(rec[0], rec[idxs[0]], col, rec[idxs[1]])
        cel(ws, "%s%d" % (col, r), f, fill,
            f_total if fill is p_destaq else f_bold, fmt=fmt, align="center")
LIN_LUCRO_TOT = r

r += 1
cel(ws, "A%d" % r, "LUCRO TOTAL (as duas receitas somadas)", p_result, f_total)
cel(ws, "B%d" % r, "=B{0}+C{0}".format(LIN_LUCRO_TOT), p_result, f_total, fmt=MOEDA, align="center")
cel(ws, "C%d" % r, None, p_result)

# ---- como usar
r += 2
ws["A%d" % r] = "COMO USAR"
ws["A%d" % r].font = f_secao
passos = [
    "1) Aba '1. Insumos': atualize o preço e o tamanho das embalagens que você compra (células amarelas).",
    "2) Abas '2. Café' e '3. Cappuccino': ajuste o rendimento, as quantidades da receita e o PREÇO DE VENDA por pote.",
    "3) Se quiser, some gás/energia/mão de obra em 'Outros custos do lote'.",
    "4) Volte aqui no Resumo para ver custo por pote, lucro e margem - e simule quantos lotes produzir.",
    "REGRA: só edite as células AMARELAS. As cinzas, verdes e laranjas são fórmulas e se atualizam sozinhas.",
]
for p in passos:
    r += 1
    ws["A%d" % r] = p
    ws["A%d" % r].font = Font(size=10, color="404040")

for col, w in zip("ABCD", (40, 22, 24, 36)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

wb.active = 0
wb.save("Planilha_Custos_Receitas.xlsx")
print("OK -> Planilha_Custos_Receitas.xlsx")
