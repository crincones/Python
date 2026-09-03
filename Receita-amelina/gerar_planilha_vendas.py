# -*- coding: utf-8 -*-
"""Gera a planilha de registro de vendas (pronta para importar no Google Planilhas)."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- estilos ----
AZUL      = "1F3864"
AZUL_CLR  = "D9E2F3"
AMARELO   = "FFF2CC"   # campo editável
CINZA     = "F2F2F2"
VERDE     = "E2EFDA"
LARANJA   = "FCE4D6"

f_titulo  = Font(name="Calibri", size=16, bold=True, color=AZUL)
f_sub     = Font(name="Calibri", size=10, italic=True, color="595959")
f_head    = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
f_secao   = Font(name="Calibri", size=11, bold=True, color=AZUL)
f_bold    = Font(name="Calibri", size=11, bold=True)
f_total   = Font(name="Calibri", size=12, bold=True, color=AZUL)

p_head    = PatternFill("solid", fgColor=AZUL)
p_edit    = PatternFill("solid", fgColor=AMARELO)
p_secao   = PatternFill("solid", fgColor=AZUL_CLR)
p_calc    = PatternFill("solid", fgColor=CINZA)
p_result  = PatternFill("solid", fgColor=VERDE)
p_destaq  = PatternFill("solid", fgColor=LARANJA)

_fino = Side(style="thin", color="BFBFBF")
b_all = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)

MOEDA = 'R$ #,##0.00'
PCT   = '0.0%'
NUM   = '#,##0.###'
DATA  = 'dd/mm/yyyy'
MES   = 'mmm/yyyy'


def cel(ws, ref, valor=None, fill=None, font=None, fmt=None, align=None, borda=True, wrap=False):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    if fill:
        c.fill = fill
    if font:
        c.font = font
    if fmt:
        c.number_format = fmt
    if align or wrap:
        c.alignment = Alignment(horizontal=align or "general", vertical="center", wrap_text=wrap)
    if borda:
        c.border = b_all
    return c


def cabecalho(ws, linha, titulos, larguras=None):
    for i, h in enumerate(titulos, start=1):
        cel(ws, "%s%d" % (get_column_letter(i), linha), h, p_head, f_head, align="center", wrap=True)
    if larguras:
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

LIN_V0, LIN_V1 = 7, 306          # linhas de lançamento na aba de vendas
LIN_P0, LIN_P1 = 6, 25           # linhas de produtos
LIN_C0, LIN_C1 = 6, 55           # linhas de clientes

PROD  = "'4. Produtos'!$A${0}:$E${1}".format(LIN_P0, LIN_P1)
PRODN = "'4. Produtos'!$A${0}:$A${1}".format(LIN_P0, LIN_P1)


def col_v(letra):
    """Faixa de uma coluna da aba de vendas, em referência absoluta."""
    return "'1. Vendas'!${0}${1}:${0}${2}".format(letra, LIN_V0, LIN_V1)


# Colunas: as que você digita ficam juntas (A a J) para caber um "colar" de uma vez;
# as calculadas ficam no fim (K a N).
V_DATA, V_CLI, V_PROD, V_QTD = col_v("A"), col_v("B"), col_v("C"), col_v("D")
V_STATUS, V_PAG, V_CANAL     = col_v("G"), col_v("H"), col_v("I")
V_TOT, V_CUSTO, V_LUCRO      = col_v("K"), col_v("L"), col_v("M")
V_MES                        = col_v("N")


# ================================================================ VENDAS =====
ws = wb.active
ws.title = "1. Vendas"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "1F3864"

ws["A1"] = "REGISTRO DE VENDAS"
ws["A1"].font = f_titulo
ws["A2"] = ("Uma linha por venda. Escolha o produto na listinha e o preço e o custo entram sozinhos "
            "(o preço pode ser trocado se você cobrou outro valor).")
ws["A2"].font = f_sub
ws["A3"] = "AMARELO = você preenche     |     CINZA = calculado automaticamente (não digite)"
ws["A3"].font = Font(size=10, bold=True, color="7F6000")
ws["A5"] = ("Para lançar uma venda nova use sempre a primeira linha vazia - as fórmulas já estão prontas "
            "até a linha %d. Colunas A a J você digita; K a N se calculam." % LIN_V1)
ws["A5"].font = f_sub

cabs_v = ["Data", "Cliente", "Produto", "Qtd", "Preço unit. (R$)", "Desconto (R$)",
          "Status", "Pagamento", "Canal", "Observações",
          "Total (R$)", "Custo (R$)", "Lucro (R$)", "Mês (auto)"]
cabecalho(ws, 6, cabs_v, [13, 24, 20, 8, 15, 14, 13, 17, 15, 28, 14, 13, 13, 13])

for r in range(LIN_V0, LIN_V1 + 1):
    vazio = '$C{0}="",$D{0}=""'.format(r)
    cel(ws, "A%d" % r, None, p_edit, fmt=DATA, align="center")
    cel(ws, "B%d" % r, None, p_edit)
    cel(ws, "C%d" % r, None, p_edit)
    cel(ws, "D%d" % r, None, p_edit, fmt=NUM, align="center")
    cel(ws, "E%d" % r, '=IF($C{0}="","",IFERROR(VLOOKUP($C{0},{1},3,FALSE),0))'.format(r, PROD),
        p_edit, fmt=MOEDA)
    cel(ws, "F%d" % r, None, p_edit, fmt=MOEDA)
    cel(ws, "G%d" % r, None, p_edit, align="center")
    cel(ws, "H%d" % r, None, p_edit, align="center")
    cel(ws, "I%d" % r, None, p_edit, align="center")
    cel(ws, "J%d" % r, None, p_edit)
    cel(ws, "K%d" % r, '=IF(OR({1}),"",$D{0}*$E{0}-N($F{0}))'.format(r, vazio),
        p_calc, f_bold, fmt=MOEDA)
    cel(ws, "L%d" % r,
        '=IF(OR({1}),"",$D{0}*IFERROR(VLOOKUP($C{0},{2},2,FALSE),0))'.format(r, vazio, PROD),
        p_calc, fmt=MOEDA)
    cel(ws, "M%d" % r, '=IF(OR({1}),"",$K{0}-$L{0})'.format(r, vazio), p_calc, f_bold, fmt=MOEDA)
    cel(ws, "N%d" % r, '=IF($A{0}="","",EOMONTH($A{0},0))'.format(r), p_calc, fmt=MES, align="center")


# ---- listinhas suspensas
def valida(coluna, faixa, titulo, msg):
    dv = DataValidation(type="list", formula1=faixa, allow_blank=True, showDropDown=False)
    dv.promptTitle = titulo
    dv.prompt = msg
    dv.errorTitle = "Opção fora da lista"
    dv.error = "Escolha uma das opções da lista (ou acrescente a sua na aba '5. Listas')."
    ws.add_data_validation(dv)
    dv.add("{0}{1}:{0}{2}".format(coluna, LIN_V0, LIN_V1))


valida("C", PRODN, "Produto", "Escolha um produto cadastrado na aba '4. Produtos'.")
valida("G", "'5. Listas'!$A$3:$A$4", "Status", "Pago ou Pendente.")
valida("H", "'5. Listas'!$B$3:$B$8", "Pagamento", "Como a cliente pagou.")
valida("I", "'5. Listas'!$C$3:$C$8", "Canal", "Por onde veio a venda.")

# ---- pendências em laranja
ws.conditional_formatting.add(
    "A{0}:N{1}".format(LIN_V0, LIN_V1),
    FormulaRule(formula=['$G{0}="Pendente"'.format(LIN_V0)], fill=p_destaq, stopIfTrue=False))

ws.auto_filter.ref = "A6:N{0}".format(LIN_V1)
ws.freeze_panes = "C7"


# ================================================================ RESUMO =====
ws = wb.create_sheet("2. Resumo")
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "1B6B57"

ws["A1"] = "RESUMO DAS VENDAS"
ws["A1"].font = f_titulo
ws["A2"] = "Esta aba só mostra resultados. Tudo vem da aba '1. Vendas' - não precisa digitar nada aqui."
ws["A2"].font = f_sub

# ---- números gerais
r = 4
cel(ws, "A%d" % r, "NÚMEROS GERAIS (todas as vendas lançadas)", p_secao, f_secao)
cel(ws, "B%d" % r, None, p_secao)

gerais = [
    ("Vendas lançadas",    '=COUNTA({0})'.format(V_PROD),  NUM,   p_calc,   f_bold),
    ("Potes vendidos",     '=SUM({0})'.format(V_QTD),      NUM,   p_calc,   f_bold),
    ("FATURAMENTO",        '=SUM({0})'.format(V_TOT),      MOEDA, p_result, f_total),
    ("Custo dos produtos", '=SUM({0})'.format(V_CUSTO),    MOEDA, p_calc,   f_bold),
    ("LUCRO",              '=SUM({0})'.format(V_LUCRO),    MOEDA, p_destaq, f_total),
]
for rot, formula, fmt, fill, font in gerais:
    r += 1
    cel(ws, "A%d" % r, rot, fill, font)
    cel(ws, "B%d" % r, formula, fill, font, fmt=fmt, align="center")

NVENDAS = "$B$%d" % (r - 4)
FAT     = "$B$%d" % (r - 2)
LUC     = "$B$%d" % r

r += 1
cel(ws, "A%d" % r, "Margem de lucro", p_destaq, f_bold)
cel(ws, "B%d" % r, '=IF({0}=0,0,{1}/{0})'.format(FAT, LUC), p_destaq, f_bold, fmt=PCT, align="center")
r += 1
cel(ws, "A%d" % r, "Ticket médio (por venda)", p_calc, f_bold)
cel(ws, "B%d" % r, '=IF({0}=0,0,{1}/{0})'.format(NVENDAS, FAT), p_calc, f_bold, fmt=MOEDA, align="center")

# ---- recebido x a receber
r += 2
cel(ws, "A%d" % r, "DINHEIRO", p_secao, f_secao)
cel(ws, "B%d" % r, None, p_secao)
for rot, formula, fill in [
        ("Já recebido (Pago)",   '=SUMIF({0},"Pago",{1})'.format(V_STATUS, V_TOT),     p_result),
        ("A RECEBER (Pendente)", '=SUMIF({0},"Pendente",{1})'.format(V_STATUS, V_TOT), p_destaq)]:
    r += 1
    cel(ws, "A%d" % r, rot, fill, f_total)
    cel(ws, "B%d" % r, formula, fill, f_total, fmt=MOEDA, align="center")

# ---- por produto
r += 2
cel(ws, "A%d" % r, "POR PRODUTO", p_secao, f_secao)
for c in "BCD":
    cel(ws, "%s%d" % (c, r), None, p_secao)
r += 1
for i, h in enumerate(["Produto", "Potes", "Faturamento (R$)", "Lucro (R$)"], 1):
    cel(ws, "%s%d" % (get_column_letter(i), r), h, p_head, f_head, align="center", wrap=True)
for k in range(LIN_P0, LIN_P1 + 1):
    r += 1
    cel(ws, "A%d" % r, "=IF('4. Produtos'!A{0}=\"\",\"\",'4. Produtos'!A{0})".format(k), p_calc, f_bold)
    cel(ws, "B%d" % r, '=IF($A{0}="","",SUMIF({1},$A{0},{2}))'.format(r, V_PROD, V_QTD),
        p_calc, fmt=NUM, align="center")
    cel(ws, "C%d" % r, '=IF($A{0}="","",SUMIF({1},$A{0},{2}))'.format(r, V_PROD, V_TOT),
        p_calc, fmt=MOEDA, align="center")
    cel(ws, "D%d" % r, '=IF($A{0}="","",SUMIF({1},$A{0},{2}))'.format(r, V_PROD, V_LUCRO),
        p_destaq, fmt=MOEDA, align="center")

# ---- por mês
r += 2
cel(ws, "A%d" % r, "MÊS A MÊS", p_secao, f_secao)
for c in "BCD":
    cel(ws, "%s%d" % (c, r), None, p_secao)
r += 1
cel(ws, "A%d" % r, "Primeiro mês do quadro", p_edit, f_bold)
cel(ws, "B%d" % r, "=EOMONTH(TODAY(),-11)", p_edit, f_bold, fmt=MES, align="center")
cel(ws, "C%d" % r, "<- edite se quiser começar em outro mês", font=f_sub, borda=False)
LIN_MES0 = r
r += 1
for i, h in enumerate(["Mês", "Potes", "Faturamento (R$)", "Lucro (R$)"], 1):
    cel(ws, "%s%d" % (get_column_letter(i), r), h, p_head, f_head, align="center", wrap=True)
for k in range(12):
    r += 1
    ref = "=$B$%d" % LIN_MES0 if k == 0 else "=EOMONTH(A%d,1)" % (r - 1)
    cel(ws, "A%d" % r, ref, p_calc, f_bold, fmt=MES, align="center")
    cel(ws, "B%d" % r, '=SUMIF({0},$A{1},{2})'.format(V_MES, r, V_QTD), p_calc, fmt=NUM, align="center")
    cel(ws, "C%d" % r, '=SUMIF({0},$A{1},{2})'.format(V_MES, r, V_TOT), p_calc, fmt=MOEDA, align="center")
    cel(ws, "D%d" % r, '=SUMIF({0},$A{1},{2})'.format(V_MES, r, V_LUCRO), p_destaq, fmt=MOEDA, align="center")

# ---- por canal / pagamento
r += 2
cel(ws, "A%d" % r, "DE ONDE VÊM AS VENDAS", p_secao, f_secao)
for c in "BCD":
    cel(ws, "%s%d" % (c, r), None, p_secao)
r += 1
for i, h in enumerate(["Canal", "Faturamento (R$)", "Forma de pagamento", "Faturamento (R$)"], 1):
    cel(ws, "%s%d" % (get_column_letter(i), r), h, p_head, f_head, align="center", wrap=True)
for k in range(6):
    r += 1
    lst = 3 + k
    cel(ws, "A%d" % r, "=IF('5. Listas'!C{0}=\"\",\"\",'5. Listas'!C{0})".format(lst), p_calc, f_bold)
    cel(ws, "B%d" % r, '=IF($A{0}="","",SUMIF({1},$A{0},{2}))'.format(r, V_CANAL, V_TOT),
        p_calc, fmt=MOEDA, align="center")
    cel(ws, "C%d" % r, "=IF('5. Listas'!B{0}=\"\",\"\",'5. Listas'!B{0})".format(lst), p_calc, f_bold)
    cel(ws, "D%d" % r, '=IF($C{0}="","",SUMIF({1},$C{0},{2}))'.format(r, V_PAG, V_TOT),
        p_calc, fmt=MOEDA, align="center")

# ---- como usar
r += 2
ws["A%d" % r] = "COMO USAR"
ws["A%d" % r].font = f_secao
passos = [
    "1) Aba '4. Produtos': confira o custo por pote (vem da planilha de custos) e o preço que você cobra.",
    "2) Aba '1. Vendas': uma linha por venda. Data, cliente, produto, quantidade e status - o resto se calcula.",
    "3) Deixe o Status em 'Pendente' enquanto não recebeu; a linha fica laranja até você trocar para 'Pago'.",
    "4) Aba '3. Clientes': escreva o nome de cada cliente uma vez para acompanhar quanto cada uma já comprou.",
    "5) Volte aqui no Resumo para ver faturamento, lucro, o que ainda tem a receber e o mês a mês.",
    "REGRA: só preencha as células AMARELAS. As cinzas, verdes e laranjas são fórmulas e se atualizam sozinhas.",
]
for p in passos:
    r += 1
    ws["A%d" % r] = p
    ws["A%d" % r].font = Font(size=10, color="404040")

for col, w in zip("ABCD", (40, 20, 26, 20)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A4"


# ============================================================== CLIENTES =====
ws = wb.create_sheet("3. Clientes")
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "A96A0C"

ws["A1"] = "CLIENTES"
ws["A1"].font = f_titulo
ws["A2"] = ("Escreva o nome de cada cliente uma vez, do mesmo jeito que você escreve na aba de vendas. "
            "Os números aparecem sozinhos.")
ws["A2"].font = f_sub

cabecalho(ws, 5, ["Cliente", "Compras", "Potes", "Total comprado (R$)",
                  "Lucro gerado (R$)", "Última compra", "Em aberto (R$)"],
          [26, 12, 10, 20, 20, 16, 16])

for r in range(LIN_C0, LIN_C1 + 1):
    cel(ws, "A%d" % r, None, p_edit)
    cel(ws, "B%d" % r, '=IF($A{0}="","",COUNTIF({1},$A{0}))'.format(r, V_CLI), p_calc, align="center")
    cel(ws, "C%d" % r, '=IF($A{0}="","",SUMIF({1},$A{0},{2}))'.format(r, V_CLI, V_QTD),
        p_calc, fmt=NUM, align="center")
    cel(ws, "D%d" % r, '=IF($A{0}="","",SUMIF({1},$A{0},{2}))'.format(r, V_CLI, V_TOT),
        p_calc, f_bold, fmt=MOEDA, align="center")
    cel(ws, "E%d" % r, '=IF($A{0}="","",SUMIF({1},$A{0},{2}))'.format(r, V_CLI, V_LUCRO),
        p_calc, fmt=MOEDA, align="center")
    cel(ws, "F%d" % r, '=IF($A{0}="","",IFERROR(MAXIFS({1},{2},$A{0}),""))'.format(r, V_DATA, V_CLI),
        p_calc, fmt=DATA, align="center")
    cel(ws, "G%d" % r,
        '=IF($A{0}="","",SUMIFS({1},{2},$A{0},{3},"Pendente"))'.format(r, V_TOT, V_CLI, V_STATUS),
        p_destaq, f_bold, fmt=MOEDA, align="center")

ws["A%d" % (LIN_C1 + 2)] = ("Escreva o nome sempre igual nas duas abas - 'Ana' e 'ana maria' contam "
                            "como clientes diferentes.")
ws["A%d" % (LIN_C1 + 2)].font = f_sub
ws.freeze_panes = "A6"


# ============================================================== PRODUTOS =====
ws = wb.create_sheet("4. Produtos")
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "833C00"

ws["A1"] = "PRODUTOS À VENDA"
ws["A1"].font = f_titulo
ws["A2"] = ("O custo por pote vem da planilha de custos (aba Resumo). "
            "Se o preço dos insumos mudar lá, atualize aqui.")
ws["A2"].font = f_sub
ws["A3"] = "AMARELO = você edita     |     CINZA = calculado automaticamente"
ws["A3"].font = Font(size=10, bold=True, color="7F6000")

cabecalho(ws, 5, ["Produto", "Custo por pote (R$)", "Preço de venda (R$)",
                  "Lucro por pote (R$)", "Margem"],
          [26, 20, 20, 20, 12])

produtos = [
    ("Café", 7.26, 25.00),
    ("Cappuccino", 9.01, 25.00),
]
for k in range(LIN_P0, LIN_P1 + 1):
    i = k - LIN_P0
    nome, custo, preco = produtos[i] if i < len(produtos) else (None, None, None)
    cel(ws, "A%d" % k, nome, p_edit, f_bold)
    cel(ws, "B%d" % k, custo, p_edit, fmt=MOEDA)
    cel(ws, "C%d" % k, preco, p_edit, fmt=MOEDA)
    cel(ws, "D%d" % k, '=IF($A{0}="","",$C{0}-$B{0})'.format(k), p_calc, f_bold, fmt=MOEDA)
    cel(ws, "E%d" % k, '=IF(OR($A{0}="",$C{0}=0),"",$D{0}/$C{0})'.format(k), p_calc, fmt=PCT, align="center")

ws["A%d" % (LIN_P1 + 2)] = ("Pode cadastrar produtos novos nas linhas vazias (kit com 2 potes, tamanho "
                            "menor, edição de Natal...) - eles aparecem na listinha da aba de vendas.")
ws["A%d" % (LIN_P1 + 2)].font = f_sub
ws.freeze_panes = "A6"


# ================================================================ LISTAS =====
ws = wb.create_sheet("5. Listas")
ws.sheet_view.showGridLines = False

ws["A1"] = "OPÇÕES DAS LISTINHAS"
ws["A1"].font = f_titulo
cabecalho(ws, 2, ["Status", "Pagamento", "Canal"], [18, 22, 22])

listas = [
    ["Pago", "Pendente"],
    ["Pix", "Dinheiro", "Cartão de débito", "Cartão de crédito", "Transferência", "Outro"],
    ["WhatsApp", "Instagram", "Feira", "Indicação", "Presente", "Outro"],
]
for c, itens in zip("ABC", listas):
    for i, item in enumerate(itens):
        cel(ws, "%s%d" % (c, 3 + i), item, p_edit)

ws["A11"] = "Pode trocar ou acrescentar opções aqui - elas aparecem nas listinhas da aba de vendas."
ws["A11"].font = f_sub

wb.active = 0
wb.save("Planilha_Vendas.xlsx")
print("OK -> Planilha_Vendas.xlsx")
